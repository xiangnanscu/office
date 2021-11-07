import os
import sys
import re
import shutil
import docx
import xlwt
import datetime

cd = os.getcwd()
sheets = {}


def delBlank(s):
    return re.sub(r'\s', '', str(s))


def getTableKey(columns):
    res = []
    for i, col in enumerate(columns):
        res.append("%s%s" % (i, delBlank(col)))
    return '|'.join(res)


for root, dirs, files in os.walk('.'):
    for f in files:
        if '~' in f or not re.search(r'docx$', f):
            continue
        docpath = os.path.abspath(os.path.join(root, f))
        print(docpath)
        doc = docx.Document(docpath)
        tables = doc.tables  # 获取文档中的所有表格
        if len(tables) == 0:
            print("no tables")
            continue
        table = tables[0]  # 提取文档中的第一个表格，测试文档中仅有一个表格
        columns = list(map(lambda e: e.text, table.row_cells(0)))
        key = getTableKey(columns)
        print(key)
        if key not in sheets:
            sheets[key] = [columns]
            print(sheets[key])
        tdata = sheets[key]
        for row in range(1, len(table.rows)):
            rowdata = []
            for col in range(0, len(table.row_cells(row))):  # 提取row行的全部列数据
                rowdata.append(delBlank(table.cell(row, col).text))
            tdata.append(rowdata)
book = xlwt.Workbook()
i = 0
for k, rows in sheets.items():
    i = i + 1
    sht = book.add_sheet(str(i))
    for r, row in enumerate(rows):
        for c, e in enumerate(row):
            sht.write(r, c, e)
book.save(os.path.join(cd, '汇总数据.xls'))
